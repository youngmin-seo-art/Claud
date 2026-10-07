import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def markdown_to_docx(md_path, docx_path, title_main):
    doc = docx.Document()
    
    # 1 inch margins (Page body width = 8.5 - 2.0 = 6.5 inches)
    page_width_inches = 6.5
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"ITS 컨버젼스 | {title_main}")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        hrun.font.name = "맑은 고딕"
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("오뉴월 코리아 프로젝트팀 | [대외비 - Confidential]")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)
        frun.font.name = "맑은 고딕"

    normal_style = doc.styles['Normal']
    normal_style.font.name = '맑은 고딕'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_lines, in_code_block
        if code_lines:
            code_text = "".join(code_lines)
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
            
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>
                    <w:top w:val="none"/>
                    <w:bottom w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(3)
            cp.paragraph_format.space_after = Pt(3)
            crun = cp.add_run(code_text.strip())
            crun.font.name = 'Consolas'
            crun.font.size = Pt(8.5)
            crun.font.color.rgb = RGBColor(15, 23, 42)
            doc.add_paragraph()
            code_lines = []
        in_code_block = False

    def flush_table():
        nonlocal table_lines, in_table
        if table_lines:
            rows_data = []
            for tline in table_lines:
                tline = tline.strip()
                if not tline.startswith('|'):
                    continue
                parts = [p.strip() for p in tline.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                rows_data.append(parts)
            
            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, "CBD5E1")

                # Calculate intelligent column widths
                col_widths = []
                if num_cols == 2:
                    col_widths = [Inches(1.8), Inches(4.7)]
                elif num_cols == 3:
                    col_widths = [Inches(1.5), Inches(2.5), Inches(2.5)]
                elif num_cols == 4:
                    col_widths = [Inches(1.2), Inches(1.8), Inches(1.8), Inches(1.7)]
                elif num_cols == 5:
                    col_widths = [Inches(1.0), Inches(1.3), Inches(1.4), Inches(1.4), Inches(1.4)]
                elif num_cols == 6:
                    col_widths = [Inches(0.9), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.2), Inches(1.1)]
                else:
                    each_w = page_width_inches / num_cols
                    col_widths = [Inches(each_w)] * num_cols

                for r_idx, r_data in enumerate(rows_data):
                    row = tbl.rows[r_idx]
                    make_row_cant_split(row)
                    is_header = (r_idx == 0)
                    if is_header:
                        make_row_header(row)

                    for c_idx in range(num_cols):
                        cell = row.cells[c_idx]
                        cell.width = col_widths[c_idx] if c_idx < len(col_widths) else Inches(1.0)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                        
                        if is_header:
                            set_cell_background(cell, "1B365D")
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, "FFFFFF")
                        else:
                            set_cell_background(cell, "F8FAFC")

                        val = r_data[c_idx] if c_idx < len(r_data) else ""
                        val_clean = val.replace('<br>', '\n').replace('•', '·')
                        
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        
                        # Parse bold inside table cells
                        parts = re.split(r'(\*\*.*?\*\*)', val_clean)
                        for part in parts:
                            if part.startswith('**') and part.endswith('**'):
                                run = p.add_run(part[2:-2])
                                run.font.bold = True
                            else:
                                run = p.add_run(part)
                            
                            run.font.name = '맑은 고딕'
                            if is_header:
                                run.font.bold = True
                                run.font.size = Pt(9.5)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                run.font.size = Pt(9)
                                run.font.color.rgb = RGBColor(30, 41, 59)

                doc.add_paragraph() # Spacing
            table_lines = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                flush_table()

        if not stripped or stripped == '---':
            continue

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:].strip())
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(27, 54, 93)
            run.font.name = '맑은 고딕'
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:].strip())
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(43, 84, 126)
            run.font.name = '맑은 고딕'
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:].strip())
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            run.font.name = '맑은 고딕'
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:].strip())
            run.font.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(51, 65, 85)
            run.font.name = '맑은 고딕'
        elif stripped.startswith('> '):
            alert_text = stripped[2:].replace('[!IMPORTANT]', '★ [중요]').replace('[!NOTE]', '■ [참고]')
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(page_width_inches)
            set_cell_background(cell, "EFF6FF")
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>
                    <w:top w:val="none"/>
                    <w:bottom w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            crun = cp.add_run(alert_text.strip())
            crun.font.size = Pt(9.5)
            crun.font.color.rgb = RGBColor(30, 58, 138)
            crun.font.name = '맑은 고딕'
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25

            content = stripped
            if content.startswith('- ') or content.startswith('* '):
                p.paragraph_format.left_indent = Inches(0.2)
                content = "• " + content[2:]
            elif re.match(r'^\d+\.\s', content):
                p.paragraph_format.left_indent = Inches(0.2)

            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    run = p.add_run(part)
                    run.font.color.rgb = RGBColor(51, 65, 85)
                run.font.name = '맑은 고딕'
                run.font.size = Pt(10)

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    # Save to primary and clean versioned path
    try:
        doc.save(docx_path)
        print(f"Successfully saved: {docx_path}")
    except PermissionError:
        alt_path = docx_path.replace(".docx", "_개정판.docx")
        doc.save(alt_path)
        print(f"Original locked by Word, saved to: {alt_path}")
        return alt_path
    return docx_path

if __name__ == "__main__":
    base_dir = r"c:\Users\immnu\Desktop\Claud\OnyuwolKorea\프로젝트 팀\국가재난_긴급_유무선_통합통신시스템_컨설팅"
    
    # 1. 요약본 변환
    sum_md = os.path.join(base_dir, "04_산출물", "요약본.md")
    sum_docx = os.path.join(base_dir, "04_산출물", "국가재난_긴급_유무선_통합통신시스템_구축컨설팅_마스터요약본.docx")
    markdown_to_docx(sum_md, sum_docx, "국가재난 긴급 유무선 통합 통신 시스템 구축 컨설팅 [요약본]")

    # 2. 본보고서 종합본 변환
    main_md = os.path.join(base_dir, "04_산출물", "본보고서_종합.md")
    main_docx = os.path.join(base_dir, "04_산출물", "국가재난_긴급_유무선_통합통신시스템_구축컨설팅_마스터본보고서_개정판.docx")
    markdown_to_docx(main_md, main_docx, "국가재난 긴급 유무선 통합 통신 시스템 구축 컨설팅 [본보고서]")
